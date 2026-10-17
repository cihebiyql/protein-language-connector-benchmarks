"""Read-only, CPU-only figure generation from the independent V3 audit.

Run with an existing Python containing pandas/matplotlib/openpyxl. To assemble
the captioned report with an existing reportlab runtime, run --pdf-only.
No model loading, fitting, imputation, p-values, or source mutations occur.
"""
# ============================================================
# CONFIGURATION — edit only this block
# ============================================================
from pathlib import Path
INPUT_FILES = [Path(__file__).resolve().parent.parent / "independent_audit.json"]
OUTPUT_DIR = Path(__file__).resolve().parent
PROJECT_TITLE = "Molecular memory V3: recovered ablation results"
CN_SERIF_FALLBACK = ["SimSun", "Songti SC", "Source Han Serif SC", "Noto Serif CJK SC", "serif"]
EN_SERIF_FALLBACK = ["Times New Roman", "Liberation Serif", "DejaVu Serif", "serif"]
OUTPUT_FORMATS = ["pdf", "png", "jpg"]
DPI = 300
ALPHA = None  # No inferential tests: at most three seeds, incomplete coverage.
SIGNIFICANCE_MARKS = {}
NS_LABEL = None
ERROR_BAR = "SD"  # Sample SD, ddof=1; only when all three planned seeds exist.
NORMALITY_ALPHA = None  # Normality testing is not informative for three seeds.
DEFAULT_DECIMALS = 3
SCIENTIFIC_LOW = 1e-3
SCIENTIFIC_HIGH = 1e5
THOUSANDS_SEP = ","
FIG_SINGLE_COL = (3.5, 2.6)
FIG_DOUBLE_COL = (8.6, 4.6)
ARMS = ["M0", "M1", "M2", "M3", "A_L", "A_Q"]
SEEDS = [17, 42, 2026]
SCENARIOS = ["new_peptide_known_target", "new_peptide_new_target"]
SCENARIO_LABELS = {SCENARIOS[0]: "New peptide / known target", SCENARIOS[1]: "New peptide / new target"}
CONTRASTS = [("Add L15", "M0", "M1"), ("Add question route", "M1", "M2"),
             ("Add cross relation", "M2", "M3"), ("Full vs no L15", "A_L", "M3"),
             ("Full vs static route", "A_Q", "M3"), ("Full vs M0", "M0", "M3")]
# ============================================================
# END CONFIGURATION
# ============================================================

import argparse
import hashlib
import json
import math
import statistics
import sys


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def report_pdf():
    """Use reportlab separately so the plotting runtime needs no new packages."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    from reportlab.platypus import SimpleDocTemplate, Image, Paragraph, Spacer, PageBreak
    from reportlab.lib.units import cm
    meta = json.loads((OUTPUT_DIR / "figure_manifest.json").read_text(encoding="utf-8"))
    title = ParagraphStyle("title", fontName="Times-Roman", fontSize=18, leading=23, alignment=TA_CENTER)
    body = ParagraphStyle("body", fontName="Times-Roman", fontSize=10, leading=15, alignment=TA_CENTER)
    caption = ParagraphStyle("caption", fontName="Times-Italic", fontSize=9, leading=12, alignment=TA_JUSTIFY)
    doc = SimpleDocTemplate(str(OUTPUT_DIR / "report.pdf"), pagesize=A4,
                            leftMargin=2.1 * cm, rightMargin=2.1 * cm,
                            topMargin=2.1 * cm, bottomMargin=2.1 * cm)
    story = [Spacer(1, 5 * cm), Paragraph(PROJECT_TITLE, title), Spacer(1, 0.7 * cm),
             Paragraph("32 of 36 planned runs accepted; four missing results remain missing.", body),
             Paragraph("Source: independent_audit.json. No model fitting or new inference.", body),
             Paragraph("Only complete 3/3 groups receive a mean and sample SD. No significance tests.", body),
             PageBreak()]
    for i, record in enumerate(meta["figures"]):
        im = Image(str(OUTPUT_DIR / record["png"]))
        aspect = im.imageHeight / im.imageWidth
        width = A4[0] - 4.2 * cm
        height = width * aspect
        if height > 17 * cm:
            height = 17 * cm
            width = height / aspect
        im.drawWidth, im.drawHeight = width, height
        story.extend([im, Spacer(1, 0.4 * cm), Paragraph(record["caption"], caption)])
        if i + 1 != len(meta["figures"]):
            story.append(PageBreak())
    doc.build(story)
    return {"report": "report.pdf", "pages": len(meta["figures"]) + 1,
            "sha256": sha256(OUTPUT_DIR / "report.pdf")}


def main():
    import importlib.metadata
    import numpy as np
    import pandas as pd
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    from matplotlib.ticker import FormatStrFormatter
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    assert OUTPUT_DIR.resolve() != INPUT_FILES[0].resolve().parent
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    process = OUTPUT_DIR / "process_data"
    process.mkdir(exist_ok=True)
    audit = json.loads(INPUT_FILES[0].read_text(encoding="utf-8-sig"))
    assert audit["verified_counts"]["canonical_accepted"] == 32
    source_sha = sha256(INPUT_FILES[0])
    records = audit["records"]
    source_keys = [r["case_id"] for r in records]
    assert len(source_keys) == len(set(source_keys)) == 36
    lookup = {}
    for r in records:
        scenario, arm, seed = r["case_id"].split("__")
        lookup[(scenario, arm, int(seed.removeprefix("seed")))] = r
    rows = []
    for scenario in SCENARIOS:
        for arm in ARMS:
            for seed in SEEDS:
                r = lookup[(scenario, arm, seed)]
                good = r["canonical_accepted"] and r["scientific_complete"]
                ev = r.get("evaluation") if good else None
                lm = ev["LM"] if good else {}
                lmreg = lm.get("regression_valid_only", {})
                num = ev["numeric"]["full_population"] if good else {}
                rows.append(dict(scenario=scenario, arm=arm, seed=seed,
                    accepted=good, queue_status=r["queue_status"],
                    n_test=ev["n_expected_test"] if good else (152 if scenario == SCENARIOS[0] else 442),
                    lm_mae=lm.get("full_sample_mae"), lm_rmse=lmreg.get("rmse"),
                    lm_spearman=lmreg.get("spearman"), lm_target_macro_mae=lmreg.get("target_macro_mae"),
                    lm_valid_rate=lm.get("parse_valid_rate"), numeric_mae=num.get("mae"),
                    numeric_rmse=num.get("rmse"), numeric_spearman=num.get("spearman"),
                    numeric_target_macro_mae=num.get("target_macro_mae")))
    df = pd.DataFrame(rows)
    metrics = [c for c in df if c.startswith("lm_") or c.startswith("numeric_")]
    assert df.shape == (36, 15), df.shape
    assert not df.duplicated(["scenario", "arm", "seed"]).any()
    assert df.accepted.sum() == 32
    assert not df.loc[df.accepted, metrics].isna().any().any()
    assert np.isfinite(df.loc[df.accepted, metrics].to_numpy()).all()
    assert df.loc[~df.accepted, metrics].isna().all().all()
    accepted = df[df.accepted]

    def md_table(frame):
        # Avoid optional tabulate dependency.
        cols = list(frame.columns)
        lines = ["| " + " | ".join(map(str, cols)) + " |", "| " + " | ".join(["---"] * len(cols)) + " |"]
        for record in frame.itertuples(index=False, name=None):
            lines.append("| " + " | ".join(str(x).replace("|", "/") for x in record) + " |")
        return "\n".join(lines)

    inventory = pd.DataFrame({"column": df.columns, "dtype": [str(df[c].dtype) for c in df],
                              "missing": [int(df[c].isna().sum()) for c in df],
                              "unique": [int(df[c].nunique()) for c in df]})
    summary = []
    for scenario in SCENARIOS:
        for arm in ARMS:
            sub = df[(df.scenario == scenario) & (df.arm == arm)]
            for channel in ["lm", "numeric"]:
                vals = sub[f"{channel}_mae"].dropna().tolist()
                summary.append(dict(scenario=scenario, arm=arm, channel=channel, k=len(vals), expected_k=3,
                    mean=statistics.mean(vals) if len(vals) == 3 else None,
                    sample_SD=statistics.stdev(vals) if len(vals) == 3 else None))
    sums = pd.DataFrame(summary)
    for row in summary:
        src = next(x for x in audit["summary"]["configurations"] if x["scenario"] == row["scenario"] and x["arm"] == row["arm"])
        metric = src["metrics"]["LM_full_sample_mae" if row["channel"] == "lm" else "numeric_mae"]
        assert row["k"] == metric["k"]
        for field in ["mean", "sample_SD"]:
            assert (row[field] is None and metric[field] is None) or math.isclose(row[field], metric[field], abs_tol=1e-12)

    outliers = []
    for metric in metrics:
        vals = accepted[metric]
        q1, q3 = vals.quantile([0.25, 0.75])
        outliers.append({"metric": metric, "IQR_flags": int(((vals < q1 - 1.5 * (q3-q1)) | (vals > q3 + 1.5 * (q3-q1))).sum()), "removed": 0})
    health = ["# Data health — recovered V3 ablations", "", f"Source SHA256: `{source_sha}`.",
        "", f"Shape: {df.shape[0]} planned rows × {df.shape[1]} columns; in-memory footprint {df.memory_usage(deep=True).sum():,} bytes.",
        "32 accepted runs, four incomplete runs. Accepted rows contain no missing or nonfinite displayed metrics. No duplicate rows or duplicate scenario/arm/seed keys.",
        "Missing accepted seeds: known-target M3 seeds 42/2026, known-target A_L seed 42, and new-target M0 seed 42. Their metrics remain missing; no partial or imputed predictions enter figures.",
        "", "## Column inventory", "", md_table(inventory), "", "## Descriptive statistics (accepted runs only)", "",
        md_table(accepted[metrics].describe().T.reset_index().rename(columns={"index": "metric"}).round(6)),
        "", "## Categories and coverage", "", md_table(sums[["scenario", "arm", "channel", "k", "expected_k"]]),
        "", "## Normality, outliers and predictors", "",
        "No normality or inferential test is justified for only three planned seeds per configuration. These are repeated model metrics, not independent predictor columns; VIF is not applicable and no regression is fitted. Correlations are descriptive and saved separately. IQR flags pooled over configurations are diagnostic only; all accepted values are retained.",
        "", md_table(pd.DataFrame(outliers)), "", "## Interpretation limits", "",
        "Only complete 3/3 groups receive a mean and sample SD (ddof=1). Paired differences use only the same accepted seed on both sides and likewise require 3/3 for summaries. No p-values, significance markers, missing-to-zero conversion, or seed selection are used. LM generation and the numeric head are separate channels. The four missing runs are execution failures, not evidence of either model success or failure. No C1 sensitivity value is treated as prediction improvement or biological causation."]
    (OUTPUT_DIR / "00_data_health.md").write_text("\n".join(health), encoding="utf-8")

    def export(frame, stem):
        frame.to_csv(process / f"{stem}.csv", index=False, float_format="%.17g")
        xlsx = process / f"{stem}.xlsx"
        frame.to_excel(xlsx, index=False, sheet_name="data")
        wb = load_workbook(xlsx)
        ws = wb.active
        ws.freeze_panes = "A2"
        for idx, col in enumerate(ws.columns, 1):
            ws.column_dimensions[get_column_letter(idx)].width = min(max(len(str(c.value or "")) for c in col) + 2, 48)
        wb.save(xlsx)

    export(df, "01_planned_and_accepted_runs")
    export(sums, "02_complete_seed_summaries")
    export(accepted[metrics].corr().reset_index().rename(columns={"index": "metric"}), "03_metric_correlations_descriptive_only")
    export(pd.DataFrame([{"analysis": "inferential_statistics", "status": "not_performed", "reason": "at most three seeds; incomplete run coverage; descriptive analysis only"}]), "05_statistical_tests")
    available = {f.name for f in font_manager.fontManager.ttflist}
    def resolve(candidates):
        return next((n for n in candidates if n in available), "serif")
    en, cn = resolve(EN_SERIF_FALLBACK), resolve(CN_SERIF_FALLBACK)
    plt.rcParams.update({"font.family": "serif", "font.serif": [en, cn, "serif"], "font.size": 10,
        "axes.labelsize": 10, "xtick.labelsize": 9, "ytick.labelsize": 9, "legend.fontsize": 9,
        "axes.unicode_minus": False, "axes.spines.top": True, "axes.spines.right": True,
        "axes.grid": True, "axes.axisbelow": True, "grid.color": "#e5e5e5", "grid.linewidth": 0.5,
        "figure.facecolor": "white", "axes.facecolor": "white", "savefig.facecolor": "white",
        "legend.frameon": False, "pdf.fonttype": 42, "savefig.dpi": DPI})
    seed_styles = {17: ("o", "#111111", -0.15), 42: ("^", "#666666", 0), 2026: ("D", "#999999", 0.15)}
    figures = []
    def save(fig, number, slug, caption):
        stem = f"fig_{number:02d}_{slug}"
        for fmt in OUTPUT_FORMATS:
            path = OUTPUT_DIR / "figures" / fmt / f"{stem}.{fmt}"
            path.parent.mkdir(parents=True, exist_ok=True)
            fig.savefig(path, dpi=DPI if fmt != "pdf" else None, bbox_inches="tight")
        plt.close(fig)
        figures.append({"stem": stem, "png": f"figures/png/{stem}.png", "caption": caption})

    for number, channel in [(1, "lm"), (2, "numeric")]:
        fig, axes = plt.subplots(1, 2, figsize=FIG_DOUBLE_COL, sharey=True)
        for ax, scenario in zip(axes, SCENARIOS):
            sub = df[df.scenario == scenario]
            for seed, (marker, color, offset) in seed_styles.items():
                sub_seed = sub[sub.seed == seed]
                ax.scatter(np.arange(len(ARMS)) + offset, sub_seed[f"{channel}_mae"], marker=marker,
                           s=30, color=color, edgecolors="black", linewidths=0.4, zorder=3, label=f"Seed {seed}")
            ns = []
            for x, arm in enumerate(ARMS):
                item = next(r for r in summary if r["scenario"] == scenario and r["arm"] == arm and r["channel"] == channel)
                ns.append(item["k"])
                if item["k"] == 3:
                    ax.errorbar(x + 0.29, item["mean"], yerr=item["sample_SD"], fmt="s", color="black", ms=4,
                                capsize=3, elinewidth=1, zorder=4, label="Mean ± sample SD" if x == 0 else None)
            ax.set_xticks(np.arange(len(ARMS)), [f"{a}\n{n}/3" for a, n in zip(ARMS, ns)])
            ax.set_xlim(-0.5, 5.6)
            ax.set_xlabel(SCENARIO_LABELS[scenario] + "\nConfiguration / accepted seeds")
            ax.set_ylabel("LM generation MAE (pKD units)" if channel == "lm" else "Numeric-head MAE (pKD units)")
            ax.yaxis.set_major_formatter(FormatStrFormatter("%.2f"))
            ax.xaxis.grid(False)
            for spine in ax.spines.values():
                spine.set_visible(True)
                spine.set_linewidth(0.8)
        vals = accepted[f"{channel}_mae"]
        complete = sums[(sums.channel == channel) & (sums.k == 3)]
        low = min(vals.min(), (complete["mean"] - complete["sample_SD"]).min())
        high = max(vals.max(), (complete["mean"] + complete["sample_SD"]).max())
        axes[0].set_ylim(max(0, low - 0.1), high + 0.15)
        handles, labels = axes[0].get_legend_handles_labels()
        fig.legend(handles, labels, loc="lower center", ncol=4, bbox_to_anchor=(0.5, -0.01))
        fig.subplots_adjust(bottom=0.27, wspace=0.13)
        export(df[["scenario", "arm", "seed", "accepted", "n_test", f"{channel}_mae"]], f"fig_{number:02d}_data")
        name = "LM generation" if channel == "lm" else "numeric-head"
        takeaway = ("Seed variation and missing full-model runs prevent a claim of stable full-architecture improvement."
                    if channel == "lm" else "The numeric head varies markedly across seeds, including errors above 3 pKD units; it must not substitute for the LM endpoint.")
        caption = (f"Figure {number}. Test {name} MAE for six V3 configurations in the two frozen scenarios "
            f"(152 and 442 test examples, respectively; lower is better). Markers identify seeds 17, 42 and 2026; "
            "black squares and error bars show mean ± sample SD only for complete 3/3 groups. "
            "Tick labels give accepted seed coverage; four unaccepted runs are omitted without imputation, "
            "and no significance tests are performed. " + takeaway)
        save(fig, number, f"{channel}_mae", caption)

    pairs = []
    for scenario in SCENARIOS:
        for label, reference, candidate in CONTRASTS:
            for channel in ["lm", "numeric"]:
                vals = []
                for seed in SEEDS:
                    a = df[(df.scenario == scenario) & (df.arm == reference) & (df.seed == seed)].iloc[0]
                    b = df[(df.scenario == scenario) & (df.arm == candidate) & (df.seed == seed)].iloc[0]
                    value = a[f"{channel}_mae"] - b[f"{channel}_mae"] if a.accepted and b.accepted else None
                    pairs.append(dict(scenario=scenario, comparison=label, reference=reference, candidate=candidate,
                                      channel=channel, seed=seed, delta_mae=value))
    paired = pd.DataFrame(pairs)
    for contrast in audit["additional_paired_contrasts"]:
        reference, candidate = contrast["reference"], contrast["candidate"]
        sub = paired[(paired.scenario == contrast["scenario"]) & (paired.reference == reference)
                     & (paired.candidate == candidate) & (paired.channel == contrast["channel"].lower())]
        assert len(sub) == 3
        for row in sub.itertuples():
            saved = contrast["per_seed"][str(row.seed)]
            assert (pd.isna(row.delta_mae) and saved is None) or math.isclose(row.delta_mae, saved, abs_tol=1e-12)
    export(paired, "fig_03_data")
    pair_summary = []
    fig, axes = plt.subplots(2, 2, figsize=(9.3, 7.3), sharex=True)
    for row, channel in enumerate(["lm", "numeric"]):
        for col, scenario in enumerate(SCENARIOS):
            ax = axes[row, col]
            ticks = []
            for j, (label, reference, candidate) in enumerate(CONTRASTS):
                subset = paired[(paired.scenario == scenario) & (paired.channel == channel) & (paired.comparison == label)]
                vals = subset.delta_mae.dropna().tolist()
                ticks.append(f"{candidate} vs {reference} ({len(vals)}/3)")
                for r in subset.itertuples():
                    if pd.notna(r.delta_mae):
                        marker, color, offset = seed_styles[r.seed]
                        ax.scatter(r.delta_mae, j + offset, color=color, marker=marker, s=28,
                                   edgecolors="black", linewidths=0.4, zorder=3)
                mean = statistics.mean(vals) if len(vals) == 3 else None
                sd = statistics.stdev(vals) if len(vals) == 3 else None
                pair_summary.append(dict(scenario=scenario, channel=channel, comparison=label, reference=reference,
                                         candidate=candidate, k=len(vals), mean=mean, sample_SD=sd,
                                         positive_count=sum(v > 0 for v in vals), negative_count=sum(v < 0 for v in vals)))
                if len(vals) == 3:
                    ax.errorbar(mean, j + 0.30, xerr=sd, fmt="s", ms=4, color="black", capsize=3, elinewidth=1)
            ax.axvline(0, color="black", ls="--", lw=0.8)
            ax.set_yticks(range(len(CONTRASTS)), ticks)
            ax.set_ylim(len(CONTRASTS) - 0.5, -0.5)
            ax.set_xlabel("Reference MAE − candidate MAE (pKD units)\nPositive: lower candidate error")
            ax.set_ylabel(("LM generation" if channel == "lm" else "Numeric head") + "\n" + SCENARIO_LABELS[scenario])
            ax.yaxis.grid(False)
            ax.xaxis.set_major_formatter(FormatStrFormatter("%.1f"))
    range_max = max(abs(paired.delta_mae.min()), abs(paired.delta_mae.max()))
    for r in pair_summary:
        if r["k"] == 3:
            range_max = max(range_max, abs(r["mean"]) + r["sample_SD"])
    axes[0, 0].set_xlim(-range_max * 1.1, range_max * 1.1)
    from matplotlib.lines import Line2D
    handles = [Line2D([0], [0], marker=mark, color="none", markerfacecolor=color,
                     markeredgecolor="black", markersize=5, label=f"Seed {seed}")
               for seed, (mark, color, _) in seed_styles.items()]
    handles.append(Line2D([0], [0], marker="s", color="black", markersize=4, label="Mean ± sample SD (3/3 only)"))
    fig.legend(handles=handles, loc="lower center", ncol=4, bbox_to_anchor=(0.5, -0.005))
    fig.subplots_adjust(bottom=0.14, left=0.20, right=0.98, hspace=0.39, wspace=0.85)
    export(pd.DataFrame(pair_summary), "04_paired_complete_seed_summaries")
    save(fig, 3, "paired_module_deltas",
         "Figure 3. Same-seed MAE differences for registered module additions and the full model in both frozen scenarios. "
         "Each marker is reference MAE minus candidate MAE for one accepted seed; positive values favour the candidate. "
         "Black squares and horizontal error bars denote mean ± sample SD only for complete 3/3 pairs, and labels show paired coverage. "
         "Seed-dependent sign changes and incomplete pairs limit claims of stable module gains; no significance tests or biological causal claims are made.")

    (OUTPUT_DIR / "captions.md").write_text("# Figure captions\n\n" + "\n\n".join(f["caption"] for f in figures), encoding="utf-8")
    manifest = {"source": "independent_audit.json", "source_sha256": source_sha,
                "execution": "CPU-only descriptive analysis; no model import, inference or fitting",
                "n_planned": 36, "n_accepted": 32, "n_missing": 4,
                "missing_policy": "retain null; summary only at 3/3; never impute",
                "inference": "none; no p-values", "error_bars": "sample SD (ddof=1), complete 3/3 groups only",
                "font": en, "dependencies": {name: importlib.metadata.version(name) for name in ["numpy", "pandas", "matplotlib", "openpyxl", "pillow"]},
                "figures": figures}
    (OUTPUT_DIR / "figure_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    # Validate exported process tables and rendering without changing sources.
    for record in figures:
        for fmt in OUTPUT_FORMATS:
            p = OUTPUT_DIR / "figures" / fmt / f"{record['stem']}.{fmt}"
            assert p.exists() and p.stat().st_size > 1000
    roundtrip = pd.read_excel(process / "01_planned_and_accepted_runs.xlsx")
    assert len(roundtrip) == 36 and roundtrip["accepted"].sum() == 32
    assert roundtrip.loc[~roundtrip["accepted"], "lm_mae"].isna().all()
    assert sha256(INPUT_FILES[0]) == source_sha, "Source audit changed while figures were built"
    print(json.dumps({"n_planned": 36, "n_accepted": 32, "n_figures": len(figures), "source_sha256": source_sha}))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pdf-only", action="store_true")
    args = parser.parse_args()
    if args.pdf_only:
        print(json.dumps(report_pdf()))
    else:
        main()
