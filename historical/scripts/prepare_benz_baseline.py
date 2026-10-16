"""Reproducible Benz EV4 engineering pilot; no unknown pair is a negative label."""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import hashlib
import json
import math
from pathlib import Path
import random
import re
import time

import numpy as np
import openpyxl

AA = "ACDEFGHIKLMNPQRSTVWY"
SEED = 42


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write_json(path, value):
    Path(path).write_text(json.dumps(value, indent=2, ensure_ascii=False, sort_keys=True) + "\n", encoding="utf-8")


class UnionFind:
    def __init__(self, n):
        self.parent = list(range(n))

    def find(self, i):
        while self.parent[i] != i:
            self.parent[i] = self.parent[self.parent[i]]
            i = self.parent[i]
        return i

    def union(self, a, b):
        a, b = self.find(a), self.find(b)
        if a != b:
            self.parent[max(a, b)] = min(a, b)


def near_duplicate(a, b, identity=0.8, coverage=0.9):
    """Existence of a gapped alignment with identity and dual residue coverage.

    Identity is matches / alignment columns (gaps count as nonmatches).
    Terminal clipping is permitted only within each sequence's coverage budget.
    DP maximizes matches - identity * columns, so a nonnegative score is an
    exact threshold witness. It does not optimize or use any target labels.
    """
    na, nb = len(a), len(b)
    ca, cb = math.ceil(na * coverage), math.ceil(nb * coverage)
    if not na or not nb:
        return False
    if sum((Counter(a) & Counter(b)).values()) + 1e-9 < identity * max(ca, cb):
        return False
    # The frequent substitution/offset cases avoid the more expensive DP.
    for shift in range(-nb + cb, na - ca + 1):
        ia, ib = max(0, shift), max(0, -shift)
        size = min(na - ia, nb - ib)
        if size >= ca and size >= cb:
            if sum(x == y for x, y in zip(a[ia:ia + size], b[ib:ib + size])) / size >= identity:
                return True
    for sa in range(na - ca + 1):
        for sb in range(nb - cb + 1):
            prev = [-identity * j for j in range(nb - sb + 1)]
            for i in range(1, na - sa + 1):
                cur = [-identity * i]
                for j in range(1, nb - sb + 1):
                    cur.append(max(prev[j - 1] + (a[sa + i - 1] == b[sb + j - 1]) - identity,
                                   prev[j] - identity, cur[-1] - identity))
                if i >= ca and max(cur[cb:]) >= -1e-9:
                    return True
                prev = cur
    return False


def expand_ev8_sequence(value):
    """Parse explicit printed alternatives, retaining any synthetic extra Y."""
    value = re.sub(r"^(H3COC|FITC)-", "", value)
    value = re.sub(r"-NH2$", "", value)
    if re.search(r"\([A-Z]+\)\([A-Z]+\)", value):
        match = re.search(r"\(([A-Z]+)\)\(([A-Z]+)\)", value)
        return [value[:match.start()] + alt + value[match.end():] for alt in match.groups()]
    match = re.search(r"\(([A-Z]+)/([A-Z]+)\)", value)
    if match:
        return [value[:match.start()] + alt + value[match.end():] for alt in match.groups()]
    if re.fullmatch(r"[" + AA + r"]+", value):
        return [value]
    raise ValueError(f"Unparsed EV8 sequence: {value}")


def read_holdout(path):
    sheet = openpyxl.load_workbook(path, read_only=True, data_only=True)["Peptides"]
    holdout = []
    for rownum, row in enumerate(sheet.values, 1):
        if rownum == 1 or not row[0]:
            continue
        name, residues, value, bait, company = row
        gene = re.sub(r"^FITC-", "", name).split(" (")[0].split("(")[0].replace("_mut", "")
        seqs = expand_ev8_sequence(value)
        start, end = map(int, residues.split("-"))
        holdout.append(dict(source="Kliche2024 EV8", row=rownum, name=name, gene=gene,
                            residues=residues, printed_sequence=value, sequences=seqs,
                            bait=bait, company=company,
                            coordinate_length_matches=all(len(s) == end - start + 1 for s in seqs),
                            quantitative_pair_membership="unresolved"))
    return holdout


def load_samples(path, targets_path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = book["ELM Benchmarking Dataset"]
    it = iter(sheet.values)
    header = next(it)
    rows = [(i, dict(zip(header, r))) for i, r in enumerate(it, 2) if r[0]]
    assert len(rows) == 3049, "Unexpected source revision: review input before changing the manifest"
    baits = sorted({r["bait"] for _, r in rows})
    target_ids = {bait: f"T{i:03d}" for i, bait in enumerate(baits, 1)}
    overview = openpyxl.load_workbook(targets_path, read_only=True, data_only=True)["EV Dataset 2b. Result overview"]
    ov = list(overview.values)
    target_meta = {r[0]: dict(zip(ov[1], r)) for r in ov[2:] if r[0]}
    targets = []
    for bait in baits:
        meta = target_meta[bait]
        targets.append({"target_id": target_ids[bait], "bait": bait,
                        **{k: meta[k] for k in ["accession", "gene_name", "protein_name", "domain_name",
                                                "domain_start", "domain_end", "pfam_id", "pfam_accession"]}})
    merged = {}
    for rownum, row in rows:
        sequence = row["Hit"].strip().upper()
        if not re.fullmatch("[" + AA + "]{1,64}", sequence):
            raise ValueError(f"Invalid sequence in source row {rownum}")
        if row["SeqStop"] - row["SeqStart"] + 1 != len(sequence):
            raise ValueError(f"Unresolved coordinates in source row {rownum}")
        sample = merged.setdefault(sequence, dict(sequence=sequence, targets=set(), parents=[], sources=[]))
        sample["targets"].add(target_ids[row["bait"]])
        gene = row["gene_name"] if isinstance(row["gene_name"], str) else None
        parent = dict(accession=row["protein_accession"], gene=gene,
                      start=row["SeqStart"], end=row["SeqStop"])
        if row["gene_name"] is not None and gene is None:
            # Excel converted some gene symbols to dates; accession remains usable.
            parent["gene_excel_corrupted_value"] = str(row["gene_name"])
        if parent not in sample["parents"]:
            sample["parents"].append(parent)
        sample["sources"].append(dict(file="Benz2022/EV4.xlsx", sheet=sheet.title, row=rownum,
                                      bait=row["bait"], confidence=row["HD2 confidence"]))
    samples = []
    for sequence, sample in sorted(merged.items()):
        sample["sample_id"] = "benz_" + hashlib.sha256(sequence.encode()).hexdigest()[:12]
        sample["targets"] = sorted(sample["targets"])
        samples.append(sample)
    return samples, targets, len(rows)


def group_samples(samples, holdout):
    anchors = [dict(sequence=seq, parents=[dict(accession=None, gene=h["gene"], start=None, end=None)],
                    holdout_row=h["row"]) for h in holdout for seq in h["sequences"]]
    nodes = samples + anchors
    uf = UnionFind(len(nodes))
    indexes = {}
    edges = Counter()
    for i, node in enumerate(nodes):
        keys = [("sequence", node["sequence"])]
        for p in node["parents"]:
            if p["accession"]:
                keys.append(("parent_accession", p["accession"].split("-")[0]))
            if p["gene"]:
                keys.append(("parent_gene", p["gene"]))
        for key in set(keys):
            if key in indexes:
                uf.union(i, indexes[key])
                edges[key[0]] += 1
            else:
                indexes[key] = i
    # Exact histogram upper bound screens impossible pairs without missing a hit.
    counts = np.array([[node["sequence"].count(aa) for aa in AA] for node in nodes], dtype=np.int16)
    lengths = np.array([len(node["sequence"]) for node in nodes])
    checked = 0
    similarity_edges = []
    for i, node in enumerate(nodes):
        overlap = np.minimum(counts[i], counts[i + 1:]).sum(axis=1)
        minimum_columns = np.maximum(np.ceil(lengths[i] * .9), np.ceil(lengths[i + 1:] * .9))
        candidates = np.nonzero(overlap + 1e-9 >= .8 * minimum_columns)[0] + i + 1
        for j in candidates:
            j = int(j)
            if uf.find(i) == uf.find(j):
                continue
            checked += 1
            if near_duplicate(node["sequence"], nodes[j]["sequence"]):
                uf.union(i, j)
                edges["identity80_dualcoverage90"] += 1
                similarity_edges.append([i, j])
    members = defaultdict(list)
    for i in range(len(nodes)):
        members[uf.find(i)].append(i)
    group_info = []
    for component in members.values():
        sample_indices = [i for i in component if i < len(samples)]
        if not sample_indices:
            continue
        ids = sorted(samples[i]["sample_id"] for i in sample_indices)
        group_id = "bg_" + hashlib.sha256("|".join(ids).encode()).hexdigest()[:12]
        anchor_rows = sorted({nodes[i]["holdout_row"] for i in component if i >= len(samples)})
        for i in sample_indices:
            samples[i]["group_id"] = group_id
            samples[i]["excluded_ev8_rows"] = anchor_rows
        group_info.append(dict(group_id=group_id, sample_ids=ids, ev8_rows=anchor_rows))
    return group_info, dict(edges), checked


def assign_splits(samples, groups, seed=SEED):
    eligible = sorted((g for g in groups if not g["ev8_rows"]), key=lambda g: g["group_id"])
    random.Random(seed).shuffle(eligible)
    ntrain = int(.7 * len(eligible))
    nval = int(.15 * len(eligible))
    for i, group in enumerate(eligible):
        group["split"] = "train" if i < ntrain else "validation" if i < ntrain + nval else "test"
    split_lookup = {g["group_id"]: g.get("split", "excluded_g1") for g in groups}
    for sample in samples:
        sample["split"] = split_lookup[sample["group_id"]]


def metrics(ranking, targets):
    target_set = set(targets)
    if not target_set:
        raise ValueError("At least one recorded positive is required")
    if len(ranking) != len(set(ranking)):
        raise ValueError("Ranking contains duplicates")
    recall = len(set(ranking[:3]) & target_set) / len(target_set)
    reciprocal = next((1 / (i + 1) for i, t in enumerate(ranking) if t in target_set), 0.0)
    return {"known_positive_recall_at_3": recall, "mrr": reciprocal}


def rank_frequency(train, target_ids):
    # One unit of weight per biological sequence, shared across its known labels.
    freq = Counter()
    for sample in train:
        for target in sample["targets"]:
            freq[target] += 1 / len(sample["targets"])
    return sorted(target_ids, key=lambda t: (-freq[t], t)), dict(freq)


def positional_similarity(a, b):
    """Maximum ungapped residue-match fraction over all offsets (B0 distance)."""
    return max(sum(x == y for x, y in zip(a[max(0, d):], b[max(0, -d):]))
               for d in range(-len(b) + 1, len(a))) / max(len(a), len(b))


def run_baselines(samples, targets):
    train = sorted([s for s in samples if s["split"] == "train"], key=lambda s: s["sample_id"])
    target_ids = [t["target_id"] for t in targets]
    frequency_order, frequency = rank_frequency(train, target_ids)
    train_matrix = np.array([[ord(c) for c in s["sequence"]] for s in train], dtype=np.uint8)
    width = train_matrix.shape[1]
    predictions = []
    for sample in samples:
        if sample["split"] not in {"validation", "test"}:
            continue
        query = np.array([ord(c) for c in sample["sequence"]], dtype=np.uint8)
        if len(query) != width:
            raise ValueError("This vectorized B0 pilot expects fixed-length Benz16 peptides")
        matches = np.zeros(len(train), dtype=np.int16)
        for offset in range(-width + 1, width):
            qa, ta = max(0, offset), max(0, -offset)
            count = min(width - qa, width - ta)
            matches = np.maximum(matches, (train_matrix[:, ta:ta + count] == query[qa:qa + count]).sum(axis=1))
        scores = matches / width
        nearest_i = int(np.argmax(scores))
        nearest = train[nearest_i]
        neighbors_order = [t for t in frequency_order if t in nearest["targets"]]
        neighbors_order += [t for t in frequency_order if t not in nearest["targets"]]
        for method, ranking in [("train_frequency", frequency_order), ("sequence_1nn", neighbors_order)]:
            predictions.append(dict(sample_id=sample["sample_id"], group_id=sample["group_id"],
                                    split=sample["split"], sequence=sample["sequence"],
                                    targets=sample["targets"], method=method, ranking=ranking,
                                    **metrics(ranking, sample["targets"]),
                                    **({"nearest_train_id": nearest["sample_id"], "nearest_similarity": scores[nearest_i]}
                                       if method == "sequence_1nn" else {})))
    summary = {}
    for split in ("validation", "test"):
        subset = [s for s in samples if s["split"] == split]
        known = {t for s in train for t in s["targets"]}
        labels = {t for s in subset for t in s["targets"]}
        counts = Counter(t for s in subset for t in s["targets"])
        summary[split] = dict(samples=len(subset), groups=len({s["group_id"] for s in subset}),
                              target_count=len(labels), targets_unseen_in_train=sorted(labels - known),
                              known_label_train_coverage=sum(t in known for s in subset for t in s["targets"])
                              / sum(len(s["targets"]) for s in subset), per_target_sample_counts=dict(counts), methods={})
        for method in ("train_frequency", "sequence_1nn"):
            pp = [p for p in predictions if p["split"] == split and p["method"] == method]
            summary[split]["methods"][method] = {k: sum(p[k] for p in pp) / len(pp)
                                                 for k in ("known_positive_recall_at_3", "mrr")}
    return predictions, summary, frequency


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()
    root = args.root
    start_time = time.perf_counter()
    raw = root / "data/raw"
    files = [raw / "benz2022/EV4.xlsx", raw / "benz2022/EV2.xlsx", raw / "kliche2024/EV8.xlsx"]
    samples, targets, nraw = load_samples(files[0], files[1])
    holdout = read_holdout(files[2])
    groups, edges, checked = group_samples(samples, holdout)
    assign_splits(samples, groups)
    processed = root / "data/processed"
    results = root / "results"
    processed.mkdir(parents=True, exist_ok=True)
    results.mkdir(parents=True, exist_ok=True)
    for name, content in [("benz_targets.json", targets), ("benz_ev8_holdout_audit.json", holdout)]:
        write_json(processed / name, content)
    with (processed / "benz_samples.jsonl").open("w", encoding="utf-8") as stream:
        for sample in samples:
            stream.write(json.dumps(sample, ensure_ascii=False) + "\n")
    manifest = dict(status="engineering_pilot_split", seed=SEED,
                    split_rule="Shuffle connected groups with seed42; floor70% train/floor15% validation/remainder test",
                    identity_threshold=.8, dual_residue_coverage=.9,
                    identity_definition="matches / gapped alignment columns; each aligned residue span >=90% full length",
                    grouping="Connected components of exact sequence, whole parent accession/gene, and threshold similarity. Whole parent is conservative over overlap-only.",
                    g1_status="All EV8 rows conservatively anchored; actual 24-pair membership, synthetic terminal residues, homologous families and cross-source D2 mapping unresolved.",
                    source_sha256={str(p.relative_to(root)): digest(p) for p in files},
                    excluded_sheets=["KPNB1", "Control Baits Dataset", "Non primary paralogue peptide m"],
                    raw_records=nraw, merged_sequences=len(samples), multilabel_samples=sum(len(s["targets"]) > 1 for s in samples),
                    target_count=len(targets), groups=groups, grouping_edge_counts=edges,
                    exact_similarity_pairs_checked=checked,
                    counts=dict(Counter(s["split"] for s in samples)))
    write_json(processed / "benz_split_manifest.json", manifest)
    print(json.dumps({"processed_ready": manifest["counts"], "groups": len(groups)}, ensure_ascii=False), flush=True)
    predictions, summary, frequency = run_baselines(samples, targets)
    summary["status"] = "engineering_pilot_not_publication_benchmark"
    summary["train_frequency_weights"] = frequency
    summary["runtime_seconds"] = time.perf_counter() - start_time
    summary["split_manifest_sha256"] = digest(processed / "benz_split_manifest.json")
    summary["evaluation"] = "Macro over unique peptide sequences; known positives only; unrecorded targets are unknown. Full34-target candidate list fixed from source; unseen training labels retain zero-frequency fallback."
    summary["nearest_neighbor"] = "Train-only 1NN by maximum ungapped match count / max length; labels of tied single best neighbor first, train-frequency fallback, target_id deterministic tie break."
    write_json(results / "b0_metrics.json", summary)
    with (results / "b0_predictions.jsonl").open("w", encoding="utf-8") as stream:
        for prediction in predictions:
            stream.write(json.dumps(prediction, ensure_ascii=False) + "\n")
    counts = manifest["counts"]
    lines = ["# B0 真实数据工程试跑", "", "本次成功跑通 Benz2022 T1 数据清洗、固定分组划分和两项训练集基线。结果为 engineering pilot，不能直接作为最终发表基准。", "",
             f"EV4 的 ELM Benchmarking Dataset 实读 {nraw} 行，合并为 {len(samples)} 条序列，其中 {manifest['multilabel_samples']} 条多标签；候选为实际出现的 {len(targets)} 个 bait。论文的 35 个研究结构域不是此表的候选数。", "",
             f"连接分组 {len(groups)} 个，实际序列数：训练 {counts.get('train',0)}，验证 {counts.get('validation',0)}，测试 {counts.get('test',0)}，EV8关联整组排除 {counts.get('excluded_g1',0)}。70/15/15 为组数比例，非精确样本比例；seed=42。", "",
             "同序列合并所有已记录阳性并保留原始行号；整母体 accession/gene 分组比仅合并重叠区段更保守。另用可含gap的序列一致性≥80%、双方残基覆盖≥90%阈值连通合组。", "",
             "EV8 全35行（含示踪肽）保守用作排除锚点，展开打印的wt/mut序列。EV8并非已确认的24对核验清单；修饰、额外Y、坐标异常和同源家族映射仍需人工对齐。D2未联合清洗，因此当前不声称完整跨源holdout清洁。", "",
             "| 划分 | 方法 | Recall@3 | MRR | 样本数 |", "|---|---|---:|---:|---:|"]
    for split in ("validation", "test"):
        for method, value in summary[split]["methods"].items():
            lines.append(f"| {split} | {method} | {value['known_positive_recall_at_3']:.6f} | {value['mrr']:.6f} | {summary[split]['samples']} |")
    lines += ["", "频率基线每条训练序列总权重为1，多标签均分；1NN仅检索训练序列，邻居标签优先后用训练频率补全排名。全部34个候选保持固定，训练未见标签不会从测试候选中删掉。指标对独立序列等权，未记录配对不作为负例。", "",
              f"测试覆盖 {summary['test']['target_count']} 个靶点，训练未见测试靶点为 {summary['test']['targets_unseen_in_train']}。逐样本完整排名及邻居ID已保存 results/b0_predictions.jsonl。", "",
             "本次验证了真实数据的清洗—划分—预测—评估路径。ESM2/Qwen接口的可学习性及连接层位比较，需进一步用B1、六组连接、T2干预、T3配对、3种子和组级置信区间评估。", "",
              "复现：`python scripts/prepare_benz_baseline.py`；回归测试：`python -m unittest discover -s tests -p test_benz_baseline.py`。输入Excel的SHA256保存在manifest；执行时间见 b0_metrics.json。"]
    (root / "reports/b0_assessment.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
