"""Preparation checks only: immutable interim input; no plot is rendered."""
import copy
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

HERE=Path(__file__).resolve().parent
SPEC=importlib.util.spec_from_file_location("paired_renderer",HERE/"render_paired_effects.py")
renderer=importlib.util.module_from_spec(SPEC);SPEC.loader.exec_module(renderer)
SOURCE=HERE.parent/"interim_20cases_20261016_2020/matched_summary.json"
SHA="4ca600c4b0d8ebf90948d98c60c2bbd2af4c3423acaf34903cafd7022f19ab3d"


class RendererPreparationTests(unittest.TestCase):
    def audit(self,source=None):
        return renderer.validate_summary(source or renderer.load_summary(SOURCE),SOURCE,SHA)

    def test_interim_grid_retains_nulls_and_no_partial_means(self):
        self.assertEqual(renderer.sha256(SOURCE),SHA)
        result=self.audit();self.assertEqual(result["errors"],[])
        self.assertEqual(result["complete_core_cases"],20)
        self.assertEqual(len(result["rows"]),30)
        self.assertEqual(sum(r["paired_difference"]is None for r in result["rows"]),5)
        for group in result["groups"]:
            if group["k"]<3:
                self.assertIsNone(group["source_mean"]);self.assertIsNone(group["source_sample_SD"])
                self.assertFalse(group["mean_displayed"])

    def test_duplicate_case_is_rejected_without_dedup(self):
        source=renderer.load_summary(SOURCE);source["cases"].append(copy.deepcopy(source["cases"][0]))
        result=self.audit(source)
        self.assertTrue(any("Duplicate case"in e for e in result["errors"]))

    def test_wrong_original_summary_is_rejected(self):
        source=renderer.load_summary(SOURCE);source["paired_contrasts"][0]["mean"]+=.1
        self.assertTrue(any("mean/sample_SD"in e for e in self.audit(source)["errors"]))

    def test_bool_numeric_string_and_nonfinite_rejected(self):
        for value in (True,"0.1",float("inf")):
            source=renderer.load_summary(SOURCE);source["paired_contrasts"][0]["per_seed"]["17"]=value
            self.assertTrue(self.audit(source)["errors"])

    def test_failure_cannot_leave_finite_full_mae(self):
        source=renderer.load_summary(SOURCE);metrics=source["cases"][0]["metrics"]
        metrics["n_valid"]-=1;metrics["n_failed"]+=1
        self.assertTrue(any("failed generation"in e for e in self.audit(source)["errors"]))

    def test_shared_limits_include_every_point_and_source_sd(self):
        audit=self.audit();lo,hi=renderer.shared_xlim(audit)
        for row in audit["rows"]:
            if renderer.finite(row["paired_difference"]):self.assertTrue(lo<row["paired_difference"]<hi)
        for group in audit["groups"]:
            if group["k"]==3:
                self.assertLess(lo,group["source_mean"]-group["source_sample_SD"])
                self.assertGreater(hi,group["source_mean"]+group["source_sample_SD"])

    def test_xlsx_exact_text_nulls_freeze_and_tuple_metadata(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory(dir=HERE)as directory:
            self.assertTrue(Path(directory).resolve().is_relative_to(HERE))
            p=Path(directory)/"test.xlsx";audit=self.audit()
            renderer.export_xlsx(p,audit,{"seed_order":renderer.SEEDS})
            book=load_workbook(p);sheet=book["plot_rows"]
            self.assertEqual(sheet.max_row,31)
            headers=[c.value for c in sheet[1]];exact=headers.index("paired_difference_exact_json")+1
            self.assertEqual(sum(sheet.cell(i,exact).value=="null"for i in range(2,32)),5)
            for row_no,row in enumerate(audit["rows"],2):
                self.assertEqual(sheet.cell(row_no,exact).value,json.dumps(row["paired_difference"]))
            self.assertTrue(all(ws.freeze_panes=="A2"for ws in book.worksheets))
            book.close()

    def test_interim_render_is_blocked_before_figures_or_xlsx(self):
        with tempfile.TemporaryDirectory(dir=HERE)as directory:
            self.assertTrue(Path(directory).resolve().is_relative_to(HERE))
            out=Path(directory)/"blocked"
            result=subprocess.run([sys.executable,str(HERE/"render_paired_effects.py"),"--summary",str(SOURCE),
                "--sha256",SHA,"--output-dir",str(out),"--render"],capture_output=True,text=True,encoding="utf-8")
            self.assertEqual(result.returncode,2,result.stderr)
            self.assertTrue((out/"00_data_health.md").exists())
            self.assertFalse((out/"figures").exists());self.assertFalse((out/"process_data").exists())
            self.assertEqual(renderer.sha256(SOURCE),SHA)


if __name__=="__main__":
    unittest.main()
