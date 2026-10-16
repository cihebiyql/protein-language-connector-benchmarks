"""Render only hash-pinned frozen summary values; no fitting or inference tests."""
# ============================================================
# CONFIGURATION — fixed prospective figure specification
# ============================================================
from pathlib import Path

INPUT_FILES = []  # Set only from the explicit --summary argument.
OUTPUT_DIR = Path(__file__).resolve().parent / "core24"
PROJECT_TITLE = "Matched connector core paired effects"
CN_SERIF_FALLBACK = ["SimSun", "Songti SC", "Source Han Serif SC", "Noto Serif CJK SC", "serif"]
EN_SERIF_FALLBACK = ["Times New Roman", "Liberation Serif", "DejaVu Serif", "serif"]
OUTPUT_FORMATS = ["pdf", "png", "jpg"]
DPI = 300
ALPHA = None
SIGNIFICANCE_MARKS = {}
NS_LABEL = None
ERROR_BAR = "SD"
NORMALITY_ALPHA = None
DEFAULT_DECIMALS = 3
SCIENTIFIC_LOW = 1e-3
SCIENTIFIC_HIGH = 1e5
THOUSANDS_SEP = ","
FIG_DOUBLE_COL = (9.2, 4.8)
SCENARIOS = ("new_peptide_known_target", "new_peptide_new_target")
SEEDS = (17, 42, 2026)
CONTRASTS = ("depth_real", "depth_neutral", "content_early", "content_mid", "DiD")
SOURCE_NAMES = {**{k: k for k in CONTRASTS[:-1]}, "DiD": "depth_content_difference_in_differences"}
FORMULAS = {"depth_real": "E_R - M_R", "depth_neutral": "E_N - M_N",
            "content_early": "E_N - E_R", "content_mid": "M_N - M_R",
            "DiD": "(E_R - M_R) - (E_N - M_N)"}
LABELS = {"depth_real": "Depth, real\nE_R - M_R", "depth_neutral": "Depth, neutral\nE_N - M_N",
          "content_early": "Content, early\nE_N - E_R", "content_mid": "Content, mid\nM_N - M_R",
          "DiD": "DiD\ndepth_real - depth_neutral"}
SEED_STYLES = {17: {"marker": "o", "color": "#555555", "offset": -.18},
               42: {"marker": "s", "color": "#888888", "offset": -.06},
               2026: {"marker": "^", "color": "#bbbbbb", "offset": .06}}
MEAN_OFFSET = .22
PROTOCOL_SHA256 = "9940d77d82e3ff23a085aa026219a4f4ee4e7ecd18bc1afe6952646a78e8b208"
EVALUATION_VERSION = "original_V1_unchanged_batch1_all36"
# ============================================================
# END CONFIGURATION
# ============================================================

import argparse
from collections import Counter
from datetime import datetime, timezone
import hashlib
import importlib.metadata
import json
import math
import os
import platform
import statistics
import sys


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def finite(value):
    return type(value) in (int, float) and math.isfinite(value)


def write_json(path, value):
    with Path(path).open("x", encoding="utf-8") as handle:
        json.dump(value, handle, ensure_ascii=False, indent=2, allow_nan=False)
        handle.write("\n")


def load_summary(path):
    def object_hook(pairs):
        if len({key for key,_ in pairs}) != len(pairs):
            raise ValueError("Duplicate JSON object key; no automatic deduplication")
        return dict(pairs)
    def reject_constant(value):
        raise ValueError("Nonstandard nonfinite JSON number: " + value)
    value=json.loads(Path(path).read_text(encoding="utf-8"),object_pairs_hook=object_hook,parse_constant=reject_constant)
    if not isinstance(value,dict):
        raise ValueError("Summary root must be an object")
    return value


def case_key(case):
    return tuple(case.get(k) for k in ("scenario", "seed", "placement", "message_input", "sender_mode", "gate_init"))


def core_key(scenario, seed, placement, content):
    return scenario, seed, placement, content, "balanced", "zero"


def numeric_close(actual, expected):
    return finite(actual) and math.isclose(actual, expected, rel_tol=1e-12, abs_tol=1e-12)


def validate_summary(summary, source_path, source_sha):
    """Check arithmetic only; retain every source null and the source summaries."""
    errors, checks = [], []
    if summary.get("schema_version") != 1:
        errors.append("Unsupported summary schema_version")
    if summary.get("protocol_sha256") != PROTOCOL_SHA256:
        errors.append("Protocol SHA differs from fixed scientific protocol")
    if summary.get("evaluation_version") != EVALUATION_VERSION:
        errors.append("Evaluation implementation differs; do not mix V1 with batching candidates")
    if summary.get("n_expected_training_trials") != 36:
        errors.append("Expected campaign is not the frozen 36-case matrix")
    cases = summary.get("cases")
    paired = summary.get("paired_contrasts")
    if not isinstance(cases, list) or not isinstance(paired, list):
        return {"errors": errors + ["cases/paired_contrasts must be lists"], "rows": [], "groups": [],
                "raw_cases": [], "checks": [], "complete_core_cases": 0, "expected_core_cases": 24,
                "extensions_not_plotted_reason": "Fixed core figure; extensions remain in original summary"}
    indexed, raw_cases, ids = {}, [], set()
    for i, item in enumerate(cases):
        c = item.get("case", {})
        key = case_key(c)
        identifier = item.get("case_id")
        if key in indexed or identifier in ids:
            errors.append(f"Duplicate case key/ID at cases[{i}]")
        if c.get("scenario") not in SCENARIOS or type(c.get("seed")) is not int or c.get("seed") not in SEEDS:
            errors.append(f"Unexpected case scenario/seed at cases[{i}]")
        indexed[key] = item
        ids.add(identifier)
        metrics = item.get("metrics")
        mae = metrics.get("full_sample_mae") if isinstance(metrics, dict) else None
        if mae is not None and not finite(mae):
            errors.append(f"Invalid full_sample_mae at cases[{i}]")
        if isinstance(metrics, dict) and item.get("status") == "complete":
            n, valid, failed = (metrics.get(k) for k in ("n_total", "n_valid", "n_failed"))
            if any(type(v) is not int or v < 0 for v in (n, valid, failed)) or valid + failed != n:
                errors.append(f"Invalid original denominator at cases[{i}]")
            elif failed > 0 and mae is not None:
                errors.append(f"A failed generation case has a defined full MAE at cases[{i}]")
            elif n != (152 if c.get("scenario") == SCENARIOS[0] else 442):
                errors.append(f"Frozen test denominator differs at cases[{i}]")
        raw_cases.append({"case_id": identifier, **c, "status": item.get("status"),
            "full_sample_mae": mae, "full_sample_mae_exact_json": json.dumps(mae),
            "n_total": metrics.get("n_total") if isinstance(metrics, dict) else None,
            "n_valid": metrics.get("n_valid") if isinstance(metrics, dict) else None,
            "n_failed": metrics.get("n_failed") if isinstance(metrics, dict) else None,
            "result_sha256": item.get("result_sha256"), "source_json_pointer": f"/cases/{i}"})
    source_complete = sum(item.get("status") == "complete" for item in cases)
    if summary.get("n_complete") != source_complete:
        errors.append("n_complete differs from case status count")
    expected_core = [core_key(s, seed, p, content) for s in SCENARIOS for seed in SEEDS
                     for p in ("early", "mid") for content in ("real", "neutral")]
    completed_core = sum(indexed.get(k, {}).get("status") == "complete" for k in expected_core)
    contrast_index = {}
    for i, group in enumerate(paired):
        key = group.get("scenario"), group.get("name")
        if key in contrast_index:
            errors.append(f"Duplicate paired contrast at index{i}")
        contrast_index[key] = (i, group)

    def constituent(scenario, seed, placement, content):
        item = indexed.get(core_key(scenario, seed, placement, content), {})
        metrics = item.get("metrics")
        value = metrics.get("full_sample_mae") if isinstance(metrics, dict) else None
        return value if item.get("status") == "complete" and finite(value) else None

    rows, groups = [], []
    for scenario in SCENARIOS:
        for contrast in CONTRASTS:
            key = scenario, SOURCE_NAMES[contrast]
            if key not in contrast_index:
                errors.append(f"Missing fixed contrast {key}")
                continue
            index, group = contrast_index[key]
            per_seed = group.get("per_seed", {})
            if not isinstance(per_seed, dict) or set(per_seed) != {str(s) for s in SEEDS}:
                errors.append(f"Expected all three explicit seed keys for {key}")
                per_seed = {}  # No inferred values are inserted into source data.
            values = [per_seed.get(str(seed)) for seed in SEEDS]
            for seed, value in zip(SEEDS, values):
                if value is not None and not finite(value):
                    errors.append(f"Nonfinite/nonnumeric/bool contrast: {key}, seed{seed}")
            k = sum(finite(v) for v in values)
            if type(group.get("k")) is not int or group.get("k") != k or group.get("expected_k") != 3:
                errors.append(f"k/expected_k mismatch for {key}")
            if group.get("coverage") != f"{k}/3":
                errors.append(f"Coverage mismatch for {key}")
            missing = [seed for seed, value in zip(SEEDS, values) if value is None]
            if set(group.get("missing_invalid_seeds", [])) != set(missing):
                errors.append(f"Missing seed metadata mismatch for {key}")
            mean, sd = group.get("mean"), group.get("sample_SD")
            if k == 3:
                if not numeric_close(mean, statistics.mean(values)) or not numeric_close(sd, statistics.stdev(values)):
                    errors.append(f"Original mean/sample_SD arithmetic mismatch for {key}")
                checks.append({"check": "original_three_seed_mean_sample_SD", "scenario": scenario,
                               "contrast": contrast, "passed": numeric_close(mean, statistics.mean(values))
                               and numeric_close(sd, statistics.stdev(values))})
            elif mean is not None or sd is not None:
                errors.append(f"Partial-seed mean/SD must remain null for {key}")
            groups.append({"scenario": scenario, "contrast": contrast, "source_contrast_name": SOURCE_NAMES[contrast],
                "k": k, "coverage": group.get("coverage"), "source_mean": mean, "source_sample_SD": sd,
                "mean_displayed": k == 3, "formula": FORMULAS[contrast],
                "source_json_pointer": f"/paired_contrasts/{index}",
                "source_mean_exact_json": json.dumps(mean), "source_sample_SD_exact_json": json.dumps(sd)})
            for seed, value in zip(SEEDS, values):
                components = {"E_R": constituent(scenario, seed, "early", "real"),
                              "E_N": constituent(scenario, seed, "early", "neutral"),
                              "M_R": constituent(scenario, seed, "mid", "real"),
                              "M_N": constituent(scenario, seed, "mid", "neutral")}
                terms = {"depth_real": ("E_R", "M_R"), "depth_neutral": ("E_N", "M_N"),
                         "content_early": ("E_N", "E_R"), "content_mid": ("M_N", "M_R")}
                needed = terms.get(contrast, ("E_R", "M_R", "E_N", "M_N"))
                if all(components[t] is not None for t in needed):
                    expected = (components[needed[0]] - components[needed[1]]) if contrast != "DiD" else (
                        (components["E_R"] - components["M_R"]) - (components["E_N"] - components["M_N"]))
                    if not numeric_close(value, expected):
                        errors.append(f"Case-MAE contrast arithmetic mismatch {key}, seed{seed}")
                elif value is not None:
                    errors.append(f"Defined contrast lacks complete finite original cases {key}, seed{seed}")
                rows.append({"scenario": scenario, "contrast": contrast, "seed": seed,
                    "paired_difference": value, "paired_difference_exact_json": json.dumps(value),
                    "paired_difference_binary64_hex": float(value).hex() if finite(value) else None,
                    "source_k": group.get("k"), "source_coverage": group.get("coverage"),
                    "source_mean": mean, "source_mean_exact_json": json.dumps(mean),
                    "source_sample_SD": sd, "source_sample_SD_exact_json": json.dumps(sd),
                    "mean_displayed": k == 3, "formula": FORMULAS[contrast],
                    **{"mae_" + name: v for name, v in components.items()},
                    "source_summary_path": str(source_path), "source_summary_sha256": source_sha,
                    "source_evaluation_version": summary.get("evaluation_version"),
                    "source_protocol_sha256": summary.get("protocol_sha256"),
                    "source_json_pointer": f"/paired_contrasts/{index}/per_seed/{seed}"})
    keys = [(r["scenario"], r["contrast"], r["seed"]) for r in rows]
    if len(rows) != 30 or len(set(keys)) != len(keys):
        errors.append("Plot grid must retain all thirty unique scenario/contrast/seed rows")
    # The two algebraic expressions must agree; no independence is assumed.
    by_key = {(r["scenario"], r["contrast"], r["seed"]): r["paired_difference"] for r in rows}
    for scenario in SCENARIOS:
        for seed in SEEDS:
            v = {c: by_key.get((scenario, c, seed)) for c in CONTRASTS}
            if all(finite(x) for x in v.values()):
                ok = numeric_close(v["DiD"], v["content_mid"] - v["content_early"]) and numeric_close(
                    v["DiD"], v["depth_real"] - v["depth_neutral"])
                if not ok:
                    errors.append(f"DiD identity mismatch {scenario}, seed{seed}")
                checks.append({"check": "DiD_algebraic_identity", "scenario": scenario, "seed": seed, "passed": ok})
    return {"errors": errors, "rows": rows, "groups": groups, "raw_cases": raw_cases, "checks": checks,
            "complete_core_cases": completed_core, "expected_core_cases": 24,
            "n_source_complete": source_complete, "source_case_records": len(cases),
            "n_original_contrasts": len(paired), "extensions_not_plotted_reason":
            "Fixed core figure; gate and pooling are prospectively registered extension stages, not score-filtered."}


def markdown_table(headers, rows):
    clean = lambda x: str(x).replace("|", "\\|").replace("\n", " ") if x is not None else "null"
    return "\n".join(["| " + " | ".join(headers) + " |", "|" + "---|" * len(headers)] +
                     ["| " + " | ".join(clean(x) for x in row) + " |" for row in rows])


def write_health(path, audit, source_path, source_sha, mode):
    import pandas as pd
    rows = audit["rows"]
    frame = pd.DataFrame(rows, dtype=object)  # Keep original None; no coercion/cleaning.
    inventory = []
    for col in frame.columns:
        values = frame[col].tolist()
        inventory.append((col, str(frame[col].dtype), json.dumps(dict(Counter(type(x).__name__ for x in values))),
                          sum(x is None for x in values), len({json.dumps(x, sort_keys=True) for x in values})))
    missing = sum(r["paired_difference"] is None for r in rows)
    lines = ["# 绘图数据健康检查", "", f"模式：{mode}；此文件先于任何图形生成。",
        f"来源：{source_path}", f"来源SHA256：`{source_sha}`", "",
        f"预期绘图网格30行；实际{len(rows)}行、{len(frame.columns)}列；DataFrame内存{int(frame.memory_usage(deep=True).sum())}字节。",
        f"核心完成{audit['complete_core_cases']}/24；有效difference {len(rows)-missing}/30，原始null {missing}/30；null保留用于coverage。",
        "没有drop、去重、清洗、插补、outlier筛选或由图形需要改变评价标准。", "", "## 列类型、缺失与唯一值", "",
        markdown_table(["column", "health dtype", "original JSON/Python types", "null", "unique"], inventory), "",
        "## 唯一键与重复", "",
        f"唯一键(scenario,contrast,seed)：{len({(r['scenario'],r['contrast'],r['seed'])for r in rows})}；重复行：{len(rows)-len({json.dumps(r,sort_keys=True)for r in rows})}。",
        "原始case/contrast重复若存在即阻止绘图，不自动去重。", "", "## 数值描述与coverage", "",
        markdown_table(["scenario", "contrast", "k/3", "finite min", "finite max", "source mean", "source sample_SD"],
            [(g["scenario"],g["contrast"],g["coverage"],
              min([r["paired_difference"] for r in rows if r["scenario"]==g["scenario"] and r["contrast"]==g["contrast"] and finite(r["paired_difference"])], default=None),
              max([r["paired_difference"] for r in rows if r["scenario"]==g["scenario"] and r["contrast"]==g["contrast"] and finite(r["paired_difference"])], default=None),
              g["source_mean"],g["source_sample_SD"])for g in audit["groups"]]), "",
        "仅核对source已有三seed均值与sample_SD的算术；少于3时不计算均值、SD或补值。没有合并不同contrast计算新的总体效应/方差。",
        "", "## 分类覆盖", "",
        markdown_table(["column", "counts"], [(key,json.dumps(dict(Counter(r[key]for r in rows)),ensure_ascii=False))
                                               for key in ("scenario","contrast","seed")]), "",
        "## 统计与解释边界", "",
        "N=3配对训练seed；DiD与其余contrast存在精确代数依赖。没有正态、VIF、相关推断、p值、优胜阈值或新CI。黑菱形误差条只可为原sample_SD，不是CI。",
        "content是分别训练的θ_R/θ_N系统总效应，不是固定θ的推理消息利用；正DiD不保证mid绝对更好。",
        audit["extensions_not_plotted_reason"], "", "## 校验结果", ""]
    lines += ["无schema/算术错误。"] if not audit["errors"] else ["绘图被阻止："] + ["- " + e for e in audit["errors"]]
    if audit["complete_core_cases"] < 24:
        lines += ["", "核心未完成24项：当前仅validation-only，禁止当作完整core24图形。"]
    Path(path).write_text("\n".join(lines)+"\n", encoding="utf-8")


def configure_process(output):
    # Changes apply only to this Python process and are made before scientific imports.
    for key in ("OMP_NUM_THREADS", "MKL_NUM_THREADS", "OPENBLAS_NUM_THREADS", "NUMEXPR_NUM_THREADS"):
        os.environ[key] = "2"
    cache = output / "matplotlib-cache"
    cache.mkdir(parents=True, exist_ok=True)
    os.environ["MPLCONFIGDIR"] = str(cache.resolve())
    os.environ["MPLBACKEND"] = "Agg"


def environment():
    packages = {}
    for name in ("matplotlib", "numpy", "pandas", "openpyxl", "Pillow", "scipy"):
        try:
            packages[name] = importlib.metadata.version(name)
        except importlib.metadata.PackageNotFoundError:
            packages[name] = None
    return {"python": sys.version, "executable": sys.executable, "platform": platform.platform(), "packages": packages,
            "process_environment": {key: os.environ.get(key) for key in
                ("OMP_NUM_THREADS","MKL_NUM_THREADS","OPENBLAS_NUM_THREADS","NUMEXPR_NUM_THREADS","MPLCONFIGDIR","MPLBACKEND")},
            "GPU_or_model_library_imported": False, "new_dependencies_installed": False}


def export_xlsx(path, audit, metadata):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    wb = Workbook(); wb.remove(wb.active)
    tables = {"plot_rows": audit["rows"], "source_summaries": audit["groups"], "raw_case_metadata": audit["raw_cases"],
              "source_metadata": [{"key": k, "value": json.dumps(v, ensure_ascii=False) if isinstance(v,(dict,list,tuple)) else v}
                                  for k,v in metadata.items()]}
    for name, records in tables.items():
        ws = wb.create_sheet(name)
        columns = list(records[0]) if records else ["no_records"]
        ws.append(columns)
        for record in records:
            ws.append([record.get(key) for key in columns])
        ws.freeze_panes = "A2"
        for n, col in enumerate(ws.columns, 1):
            ws.column_dimensions[get_column_letter(n)].width = min(50,max(12,max(len(str(c.value))if c.value is not None else 0 for c in col)+2))
    wb.save(path)


def shared_xlim(audit):
    values = [r["paired_difference"] for r in audit["rows"] if finite(r["paired_difference"])] + [0.]
    for g in audit["groups"]:
        if g["k"] == 3:
            values += [g["source_mean"] - g["source_sample_SD"], g["source_mean"] + g["source_sample_SD"]]
    lo, hi = min(values), max(values)
    width = hi-lo if hi>lo else 1.
    return lo-.08*width, hi+.08*width


def make_caption(audit):
    summaries = {(g["scenario"],g["contrast"]):g for g in audit["groups"]}
    descriptions=[]
    for scenario,label in zip(SCENARIOS,("已见靶点","双新")):
        g=summaries[(scenario,"depth_real")]
        descriptions.append(f"{label}depth_real的原三seed均值为{g['source_mean']:.4f}" if g["k"]==3
                            else f"{label}depth_real覆盖{g['coverage']}，不报告跨seed均值")
    return ("图1。两个冻结A1场景下五项预先固定核心配对contrast：depth_real、depth_neutral、content_early、content_mid与DiD，横轴为MAE差（pKD）。\n\n"
        "灰度圆、方、三角分别表示同一seed下两系统的配对差值（17、42、2026）；仅k＝3时以黑菱形和误差条显示原summary均值±sample_SD，误差条为seed间SD而非CI。原null保留coverage，未插补、删seed或实施推断检验；冻结test测量不因seed重复变成新增独立生物样本。\n\n"
        +"；".join(descriptions)+"。深度正值表示相应内容条件下mid的MAE更低；content是独立训练θ_R与θ_N系统的总效应，不是固定θ推理消息使用。DiD由其他contrast代数决定，正值不保证mid绝对更好；门控和池化属于预登记扩展阶段，因此未纳入此核心主图，原summary保留全部记录。\n")


def draw(audit, output):
    import matplotlib
    matplotlib.use("Agg", force=True)
    from matplotlib import font_manager
    import matplotlib.pyplot as plt
    from matplotlib.lines import Line2D
    from matplotlib.ticker import FuncFormatter, MaxNLocator
    available={f.name for f in font_manager.fontManager.ttflist}
    resolve=lambda candidates:next((f for f in candidates if f in available),"serif")
    en,cn=resolve(EN_SERIF_FALLBACK),resolve(CN_SERIF_FALLBACK)
    matplotlib.rcParams.update({"font.family":"serif","font.serif":[en,cn,"DejaVu Serif"],"font.size":10,
        "axes.labelsize":10,"xtick.labelsize":9,"ytick.labelsize":9,"legend.fontsize":9,
        "axes.unicode_minus":False,"axes.edgecolor":"black","axes.labelcolor":"black",
        "figure.facecolor":"white","axes.facecolor":"white","savefig.facecolor":"white",
        "legend.frameon":False,"pdf.fonttype":42,"ps.fonttype":42})
    fig,axes=plt.subplots(1,2,figsize=FIG_DOUBLE_COL,sharex=True,sharey=True)
    bounds=shared_xlim(audit)
    groups={(g["scenario"],g["contrast"]):g for g in audit["groups"]}
    def number(x,_pos=None):
        if x==0:return "0"
        return f"{x:.{DEFAULT_DECIMALS}e}" if abs(x)<SCIENTIFIC_LOW or abs(x)>=SCIENTIFIC_HIGH else f"{x:.{DEFAULT_DECIMALS}g}"
    for ax,scenario,panel in zip(axes,SCENARIOS,("(a) New peptide / known target","(b) New peptide / new target")):
        for y,contrast in enumerate(CONTRASTS):
            for seed in SEEDS:
                row=next(r for r in audit["rows"]if(r["scenario"],r["contrast"],r["seed"])==(scenario,contrast,seed))
                if finite(row["paired_difference"]):
                    style=SEED_STYLES[seed]
                    ax.plot(row["paired_difference"],y+style["offset"],marker=style["marker"],markersize=5,
                            color=style["color"],markeredgecolor="black",markeredgewidth=.45,linestyle="none",zorder=3)
            g=groups[(scenario,contrast)]
            if g["k"]==3:
                ax.errorbar(g["source_mean"],y+MEAN_OFFSET,xerr=g["source_sample_SD"],fmt="D",markersize=5,
                            color="black",ecolor="black",capsize=3,elinewidth=.9,zorder=4)
            else:
                ax.text(.98,y+MEAN_OFFSET,f"k={g['coverage']}; no mean",transform=ax.get_yaxis_transform(),
                        ha="right",va="center",fontsize=8,color="#444444")
        ax.axvline(0,color="#444444",linestyle="--",linewidth=.8,zorder=1)
        ax.set_xlim(*bounds);ax.set_ylim(4.55,-.55)
        ax.set_yticks(range(5),[LABELS[c]for c in CONTRASTS])
        ax.set_xlabel("MAE difference (pKD)")
        ax.xaxis.set_major_formatter(FuncFormatter(number));ax.xaxis.set_major_locator(MaxNLocator(nbins=5,min_n_ticks=3))
        ax.xaxis.grid(True,color="#e5e5e5",linewidth=.5);ax.yaxis.grid(False);ax.set_axisbelow(True)
        for side in ("top","right","left","bottom"):
            ax.spines[side].set_visible(True);ax.spines[side].set_color("black");ax.spines[side].set_linewidth(.8)
        ax.text(0,1.025,panel,transform=ax.transAxes,ha="left",va="bottom",fontsize=10)
    handles=[Line2D([],[],marker=SEED_STYLES[s]["marker"],color=SEED_STYLES[s]["color"],
                    markeredgecolor="black",markeredgewidth=.45,linestyle="none",label=f"Seed {s}",markersize=5)for s in SEEDS]
    handles.append(Line2D([],[],marker="D",color="black",linestyle="-",label="Mean ± sample SD (k=3)",markersize=5))
    fig.legend(handles=handles,loc="lower center",bbox_to_anchor=(.53,.005),ncol=4)
    fig.subplots_adjust(left=.22,right=.98,bottom=.17,top=.91,wspace=.12)
    files=[]
    for fmt in OUTPUT_FORMATS:
        path=output/"figures"/fmt/f"fig_01_core_paired_effects.{fmt}";path.parent.mkdir(parents=True,exist_ok=True)
        kwargs={"dpi":DPI,"bbox_inches":"tight","pad_inches":.10}
        if fmt=="jpg":kwargs["pil_kwargs"]={"quality":95,"subsampling":0}
        fig.savefig(path,**kwargs);files.append(path)
    plt.close(fig)
    return files,{"english_serif":en,"chinese_serif":cn,"backend":matplotlib.get_backend(),"shared_xlim":bounds,
                  "means_and_SD_from_original_source":True,"seed_markers_offsets":SEED_STYLES,"mean_offset":MEAN_OFFSET}


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--summary",type=Path,required=True)
    parser.add_argument("--sha256",required=True)
    parser.add_argument("--output-dir",type=Path,default=OUTPUT_DIR)
    mode=parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--validate-only",action="store_true")
    mode.add_argument("--render",action="store_true")
    args=parser.parse_args()
    source=args.summary.resolve();output=args.output_dir.resolve()
    if output==source.parent or output==source or source.is_relative_to(output):
        raise ValueError("Outputs must be separate from the immutable input snapshot")
    if sha256(source)!=args.sha256:
        raise ValueError("Source SHA mismatch; source not read as accepted scientific data")
    if output.exists() and any(output.iterdir()):
        raise ValueError("Use a new output directory; historical validation/figure artifacts are never overwritten")
    output.mkdir(parents=True,exist_ok=True)
    configure_process(output)
    try:
        summary=load_summary(source)
    except ValueError as error:
        (output/"00_data_health.md").write_text("# 绘图数据健康检查\n\n来源SHA：`"+args.sha256+
            "`\n\n源JSON校验失败；未修复、清洗或绘图："+str(error)+"\n",encoding="utf-8")
        write_json(output/"validation_receipt.json",{"accepted_schema_and_arithmetic":False,
            "source_summary_sha256":args.sha256,"errors":[str(error)],"figures_created":0})
        return 2
    audit=validate_summary(summary,source,args.sha256)
    if args.render and audit["complete_core_cases"]!=24:
        audit["errors"].append("Rendering requires all24 completed immutable core cases; interim input is validation-only")
    write_health(output/"00_data_health.md",audit,source,args.sha256,"render"if args.render else"validate-only")
    metadata={"source_summary_path":str(source),"source_summary_sha256":args.sha256,
              "renderer_sha256":sha256(__file__),"evaluation_version":summary.get("evaluation_version"),
              "protocol_sha256":summary.get("protocol_sha256"),"source_n_complete":summary.get("n_complete"),
              "core_complete":audit["complete_core_cases"],"plot_rows":len(audit["rows"]),
              "fixed_contrasts":CONTRASTS,"seed_order":SEEDS,"uncertainty":"original sample_SD of paired seed effects; not CI",
              "no_cleaning_imputation_filtering":True,"inference_tests_performed":False,
              "extension_stage_exclusion":audit["extensions_not_plotted_reason"]}
    write_json(output/"validation_receipt.json",{"mode":"render"if args.render else"validate-only",**metadata,
        "accepted_schema_and_arithmetic":not audit["errors"],"errors":audit["errors"],"checks":audit["checks"],
        "plot_rows_including_original_nulls":audit["rows"],"source_groups":audit["groups"],"environment":environment()})
    if audit["errors"]:
        print(json.dumps({"status":"blocked","errors":audit["errors"],"data_health":str(output/"00_data_health.md")},ensure_ascii=False))
        return 2
    if args.validate_only:
        assert sha256(source)==args.sha256
        print(json.dumps({"status":"validated_only","core_complete":audit["complete_core_cases"],"rows":len(audit["rows"]),"figures_created":0},ensure_ascii=False))
        return 0
    process=output/"process_data";process.mkdir()
    workbook=process/"fig_01_data.xlsx";export_xlsx(workbook,audit,metadata)
    caption=output/"fig_01_caption_zh.md";caption.write_text(make_caption(audit),encoding="utf-8")
    figures,style=draw(audit,output)
    assert sha256(source)==args.sha256,"Immutable source changed during rendering"
    outputs=[output/"00_data_health.md",output/"validation_receipt.json",workbook,caption,*figures]
    write_json(output/"manifest.json",{"created_utc":datetime.now(timezone.utc).isoformat(),**metadata,
        "environment":environment(),"style":style,"output_sha256":{str(p.relative_to(output)):sha256(p)for p in outputs},
        "merged_report_pdf_created":False,"meaning":"descriptive fixed-core figure; no new inference statistics"})
    print(json.dumps({"status":"rendered","output":str(output),"figures":len(figures),"plot_rows":30},ensure_ascii=False))
    return 0


if __name__=="__main__":
    raise SystemExit(main())
